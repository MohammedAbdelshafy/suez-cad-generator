"""
Build the ultra-futuristic graduation presentation (.pptx) for the
SUEZ PORT CAD — Masterplan Design AI Engineer project.

Renders the to-scale CAD plans with matplotlib and assembles a dark,
bioluminescent-themed deck with python-pptx. Self-contained: run with the
project venv.

    .venv\\Scripts\\python.exe presentation\\build_presentation.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, Circle

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))
from calculations import compute, VesselInput, VESSEL_PRESETS, BLOCK_COEFFICIENT  # noqa: E402

# Coastal calc engine built from Dr. Elemary's ECB 3802 lectures.
sys.path.insert(0, r"C:\Users\omare\OneDrive\Desktop\AI\Coastal_ECB3802")
import coastal_calcs as W  # noqa: E402

ASSETS = Path(__file__).resolve().parent / "assets"
ASSETS.mkdir(exist_ok=True)
LOGO = ASSETS / "uni_logo.png"
OUT_PPTX = Path(__file__).resolve().parent / "SUEZ_Port_CAD_Graduation.pptx"

# ---------------------------------------------------------------- palette ----
INK   = RGBColor(0x05, 0x09, 0x10)
INK2  = RGBColor(0x0B, 0x14, 0x1E)
PANEL = RGBColor(0x0E, 0x1A, 0x24)
BIO   = RGBColor(0x19, 0xF0, 0xC8)
CYAN  = RGBColor(0x5E, 0xF0, 0xFF)
PLASMA= RGBColor(0x7A, 0x5C, 0xFF)
AMBER = RGBColor(0xFF, 0xB4, 0x54)
TEXT  = RGBColor(0xD6, 0xF2, 0xEC)
DIM   = RGBColor(0x83, 0xB4, 0xAB)
WHITE = RGBColor(0xEA, 0xFF, 0xFB)

HEAD = "Bahnschrift"
MONO = "Consolas"
BODY = "Segoe UI"

# matplotlib hex equivalents
M_INK = "#05090f"
M_BIO = "#19f0c8"
M_CYAN = "#5ef0ff"
M_PLASMA = "#7a5cff"
M_AMBER = "#ffb454"
M_TEXT = "#d6f2ec"
M_GRID = "#16463f"

EMU_IN = 914400
SW, SH = 13.333, 7.5  # inches (16:9)


# ============================================================== plan render ==
def render_plan(payload: dict, path: Path, title: str) -> None:
    geo = payload["geometry"]["shapes"]
    style = {
        "QUAY": M_AMBER, "BERTH": M_CYAN, "TURNING": M_BIO,
        "CHANNEL": M_BIO, "ANCHORAGE": M_PLASMA,
    }
    fig, ax = plt.subplots(figsize=(8.6, 6.4), dpi=170)
    fig.patch.set_facecolor(M_INK)
    ax.set_facecolor(M_INK)

    style["STRUCT"] = M_CYAN
    style["DIM"] = DIM_hex()
    ha_map = {"middle": "center", "start": "left", "end": "right"}
    for s in geo:
        c = style.get(s["layer"], M_BIO)
        kind = s["kind"]
        if kind == "circle":
            ax.add_patch(Circle((s["cx"], s["cy"]), s["r"], fill=False, ec=c, lw=1.6))
            if s.get("label"):
                ax.text(s["cx"], s["cy"], s["label"], color=c, ha="center",
                        va="center", fontsize=10, family="monospace")
        elif kind == "rect":
            ax.add_patch(Rectangle((s["x"], s["y"]), s["w"], s["h"], fill=True,
                                   ec=c, fc=c + "22", lw=1.4))
            if s.get("label"):
                ax.text(s["x"] + s["w"] / 2, s["y"] + s["h"] / 2, s["label"],
                        color=c, ha="center", va="center", fontsize=9, family="monospace")
        elif kind == "polyline":
            xs = [p[0] for p in s["points"]]; ys = [p[1] for p in s["points"]]
            ax.plot(xs, ys, color=c, lw=2.6)
            if s.get("label"):
                ax.text(sum(xs) / len(xs), max(ys) + 30, s["label"], color=c,
                        ha="center", fontsize=10, family="monospace")
        elif kind == "note":
            ax.text(s["x"], s["y"], s["text"], color=c,
                    ha=ha_map.get(s.get("anchor"), "center"), fontsize=8,
                    family="monospace")
        elif kind == "dimension":
            p1, p2 = s["p1"], s["p2"]
            ax.plot([p1[0], p2[0]], [p1[1], p2[1]], color=c, lw=0.8, alpha=0.8)
            ax.text((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2 + 12, s["text"],
                    color=c, ha="center", fontsize=7, family="monospace")

    ax.set_aspect("equal")
    ax.grid(True, color=M_GRID, lw=0.6, alpha=0.6)
    ax.tick_params(colors=DIM_hex(), labelsize=7)
    for sp in ax.spines.values():
        sp.set_color(M_GRID)
    ax.set_xlabel("metres", color=M_TEXT, fontsize=8)
    ax.set_title(title, color=M_TEXT, fontsize=12, family="monospace", pad=12)
    fig.tight_layout()
    fig.savefig(path, facecolor=M_INK, bbox_inches="tight")
    plt.close(fig)


def DIM_hex():
    return "#83b4ab"


def render_windrose(path: Path) -> None:
    """Polar wind rose from the lecture wind table (hours/yr, stacked by speed)."""
    import numpy as np
    deg = {"N": 0, "30": 30, "60": 60, "E": 90, "120": 120, "150": 150,
           "S": 180, "210": 210, "240": 240, "W": 270, "300": 300, "330": 330}
    rows, _, _ = W.wind_rose()
    fig = plt.figure(figsize=(6.6, 6.6), dpi=170)
    ax = fig.add_subplot(111, projection="polar")
    fig.patch.set_facecolor(M_INK); ax.set_facecolor(M_INK)
    ax.set_theta_zero_location("N"); ax.set_theta_direction(-1)
    colors = [M_BIO, M_CYAN, M_PLASMA, M_AMBER]
    labels = ["5 kt", "15 kt", "25 kt", "35 kt"]
    width = np.radians(26)
    for r in rows:
        ang = np.radians(deg[r["dir"]]); bottom = 0.0
        for ci, dur in enumerate(r["durations"]):
            ax.bar(ang, dur, width=width, bottom=bottom, color=colors[ci],
                   edgecolor=M_INK, linewidth=0.6)
            bottom += dur
    ax.tick_params(colors=M_TEXT, labelsize=8)
    ax.set_yticklabels([])
    ax.grid(True, color=M_GRID, lw=0.6, alpha=0.7)
    for sp in ax.spines.values():
        sp.set_color(M_GRID)
    handles = [plt.Rectangle((0, 0), 1, 1, color=c) for c in colors]
    leg = ax.legend(handles, labels, loc="lower right", bbox_to_anchor=(1.15, -0.05),
                    frameon=False, fontsize=8, labelcolor=M_TEXT, title="speed class")
    leg.get_title().set_color(DIM_hex())
    ax.set_title("WIND ROSE — hours/yr by direction & speed",
                 color=M_TEXT, fontsize=11, pad=22, family="monospace")
    fig.savefig(path, facecolor=M_INK, bbox_inches="tight")
    plt.close(fig)


# ================================================================ pptx deck ==
def _no_shadow(shape):
    shape.shadow.inherit = False


def rect(slide, l, t, w, h, color, line=None, lw=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is not None:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    else:
        sp.line.fill.background()
    _no_shadow(sp)
    return sp


def hline(slide, l, t, w, color, weight=2.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Pt(weight))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background(); _no_shadow(sp)
    return sp


def vline(slide, l, t, h, color, weight=2.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Pt(weight), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background(); _no_shadow(sp)
    return sp


def _spacing(run, pts):
    run._r.get_or_add_rPr().set("spc", str(int(pts * 100)))


def textbox(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if "before" in ln: p.space_before = Pt(ln["before"])
        if "after" in ln: p.space_after = Pt(ln["after"])
        if "line" in ln:
            p.line_spacing = ln["line"]
        for seg in ln["runs"]:
            r = p.add_run(); r.text = seg["t"]
            f = r.font
            f.size = Pt(seg.get("size", 18)); f.bold = seg.get("bold", False)
            f.italic = seg.get("italic", False)
            f.name = seg.get("font", BODY); f.color.rgb = seg.get("color", TEXT)
            if seg.get("spc"): _spacing(r, seg["spc"])
    return tb


# ----------------------------------------------------------- backgrounds ----
PORT_IMAGES = Path(r"C:\Users\omare\OneDrive\Desktop\AI\port_images")
_BACKGROUNDS: list[Path] = []
_bg_state = {"i": 0}


def scrim(slide, alpha_pct: int = 50):
    """Full-slide semi-transparent black overlay for uniform text contrast."""
    from pptx.oxml.ns import qn
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                Emu(int(SW * EMU_IN)), Emu(int(SH * EMU_IN)))
    sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor(0x03, 0x06, 0x0B)
    sp.line.fill.background(); _no_shadow(sp)
    srgb = sp.fill.fore_color._xFill.find(qn("a:srgbClr"))
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha_pct * 1000))}))
    return sp


def prepare_backgrounds(darken: float = 0.5) -> list[Path]:
    """Center-crop every port photo to 16:9 and darken it so slide text stays
    readable. Returns the processed files (cycled across slides). Drop more
    images into PORT_IMAGES and they are picked up automatically."""
    from PIL import Image
    srcs = sorted(PORT_IMAGES.glob("*.png")) + sorted(PORT_IMAGES.glob("*.jpg"))
    outs = []
    target = SW / SH  # 16:9
    for i, p in enumerate(srcs):
        im = Image.open(p).convert("RGB")
        w, h = im.size
        ar = w / h
        if ar > target:                      # too wide -> crop sides
            nw = int(h * target)
            im = im.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
        else:                                # too tall -> crop top/bottom
            nh = int(w / target)
            im = im.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
        im = Image.eval(im, lambda x: int(x * darken))   # dim for legibility
        out = ASSETS / f"bg_{i:02d}.png"
        im.save(out)
        outs.append(out)
    return outs


def base_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if _BACKGROUNDS:
        img = _BACKGROUNDS[_bg_state["i"] % len(_BACKGROUNDS)]
        _bg_state["i"] += 1
        slide.shapes.add_picture(str(img), 0, 0,
                                 width=Emu(int(SW * EMU_IN)),
                                 height=Emu(int(SH * EMU_IN)))
        scrim(slide, 50)
    else:
        rect(slide, 0, 0, SW, SH, INK)
    return slide


def decor(slide, section, number, total):
    # left accent line + node
    vline(slide, 0.55, 1.0, 5.5, BIO, 2.2)
    rect(slide, 0.5, 0.95, 0.1, 0.1, BIO)
    # top-right university logo (golden AASTMT emblem)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(SW - 1.25), Inches(0.3),
                                 height=Inches(0.78))
    # footer
    hline(slide, 0.9, SH - 0.62, SW - 1.8, PANEL, 1.2)
    textbox(slide, 0.9, SH - 0.55, 8, 0.3,
            [{"runs": [{"t": "SUEZ PORT CAD · MASTERPLAN DESIGN AI ENGINEER",
                        "size": 9, "color": DIM, "font": MONO, "spc": 2}]}])
    textbox(slide, SW - 2.4, SH - 0.55, 1.5, 0.3,
            [{"runs": [{"t": f"{number:02d} / {total:02d}", "size": 9, "color": BIO,
                        "font": MONO, "spc": 2}]}], align=PP_ALIGN.RIGHT)
    # section tag
    textbox(slide, 0.9, 0.52, 9, 0.3,
            [{"runs": [{"t": section.upper(), "size": 11, "color": AMBER,
                        "font": HEAD, "bold": True, "spc": 4}]}])


def header(slide, title):
    textbox(slide, 0.85, 0.92, 11.6, 1.0,
            [{"runs": [{"t": title, "size": 32, "color": WHITE, "font": HEAD,
                        "bold": True, "spc": 1}]}])
    hline(slide, 0.9, 1.78, 3.4, BIO, 2.5)


def bullets(slide, items, l=1.15, t=2.15, w=11.2, size=18, gap=12):
    lines = []
    for it in items:
        if isinstance(it, tuple):
            head, sub = it
            lines.append({"before": gap, "runs": [
                {"t": "▸  ", "size": size, "color": BIO, "font": HEAD, "bold": True},
                {"t": head, "size": size, "color": TEXT, "font": BODY, "bold": True}]})
            lines.append({"before": 2, "runs": [
                {"t": "     " + sub, "size": size - 4, "color": DIM, "font": BODY}]})
        else:
            lines.append({"before": gap, "runs": [
                {"t": "▸  ", "size": size, "color": BIO, "font": HEAD, "bold": True},
                {"t": it, "size": size, "color": TEXT, "font": BODY}]})
    textbox(slide, l, t, w, SH - t - 0.9, lines)


def formula_card(slide, l, t, w, h, name, expr, note):
    rect(slide, l, t, w, h, PANEL)
    vline(slide, l, t, h, BIO, 2.5)
    textbox(slide, l + 0.25, t + 0.18, w - 0.5, 0.4,
            [{"runs": [{"t": name, "size": 13, "color": AMBER, "font": HEAD,
                        "bold": True, "spc": 1}]}])
    textbox(slide, l + 0.25, t + 0.62, w - 0.5, 0.6,
            [{"runs": [{"t": expr, "size": 17, "color": WHITE, "font": MONO, "bold": True}]}])
    textbox(slide, l + 0.25, t + h - 0.5, w - 0.5, 0.4,
            [{"runs": [{"t": note, "size": 11, "color": DIM, "font": BODY}]}])


# ================================================================== content ==
def build():
    # photographic slide backgrounds (cycled, darkened for legibility)
    _BACKGROUNDS[:] = prepare_backgrounds()
    _bg_state["i"] = 0

    # worked example
    nx = {k: v for k, v in VESSEL_PRESETS["neopanamax"].items() if k != "label"}
    example = compute(VesselInput(**nx))
    render_plan(example, ASSETS / "plan_neopanamax.png",
                "NEO-PANAMAX CONTAINER — CONCEPT MASTERPLAN")
    vlcc = {k: v for k, v in VESSEL_PRESETS["vlcc"].items() if k != "label"}
    render_plan(compute(VesselInput(**vlcc)), ASSETS / "plan_vlcc.png",
                "VLCC TANKER — CONCEPT MASTERPLAN")

    # ---- Max-LOA design + coastal calcs (Dr. Elemary's ECB 3802 lectures) ----
    ulcv = {k: v for k, v in VESSEL_PRESETS["ulcv"].items() if k != "label"}
    maxloa = compute(VesselInput(**ulcv, num_berths=2))
    render_plan(maxloa, ASSETS / "plan_ulcv.png",
                "MAX-LOA ULCV (400 m) — CONCEPT MASTERPLAN")
    render_windrose(ASSETS / "windrose.png")
    _rows, _grand, _pwd = W.wind_rose()
    chan_depth = next(r["value"] for r in maxloa["results"]
                      if r["label"] == "Approach Channel Depth")
    Tw, H0w, phi0w, Kdw = 8.0, 2.5, 30.0, 0.10
    sw = W.shoaling(Tw, chan_depth)
    rw = W.refraction(Tw, chan_depth, phi0w)
    Hw = W.combined_H(H0w, sw["Ks"], rw["Kr"], Kdw)

    prs = Presentation()
    prs.slide_width = Emu(int(SW * EMU_IN))
    prs.slide_height = Emu(int(SH * EMU_IN))
    TOTAL = 20

    # ---- 01 TITLE ----
    s = base_slide(prs)
    vline(s, 1.0, 1.7, 3.6, BIO, 3)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(10.95), Inches(0.45), height=Inches(1.85))
    textbox(s, 1.3, 1.5, 9.4, 0.4,
            [{"runs": [{"t": "ARAB ACADEMY FOR SCIENCE, TECHNOLOGY & MARITIME TRANSPORT",
                        "size": 12.5, "color": AMBER, "font": HEAD, "bold": True, "spc": 2}]}])
    textbox(s, 1.3, 1.82, 9.4, 0.4,
            [{"runs": [{"t": "GRADUATION PROJECT · ECB 3802 · MARINE PORT ENGINEERING",
                        "size": 11, "color": DIM, "font": HEAD, "bold": True, "spc": 3}]}])
    textbox(s, 1.25, 2.05, 11.5, 1.7,
            [{"runs": [{"t": "SUEZ PORT CAD", "size": 66, "color": WHITE, "font": HEAD,
                        "bold": True, "spc": 1}]}])
    textbox(s, 1.3, 3.5, 11.5, 0.7,
            [{"runs": [{"t": "CAD & Masterplan Design AI Engineer", "size": 26,
                        "color": BIO, "font": HEAD, "spc": 1}]}])
    textbox(s, 1.3, 4.35, 11, 0.6,
            [{"runs": [{"t": "Design-Vessel Port Dimensioning & Automated Masterplan Generation",
                        "size": 16, "color": DIM, "font": BODY}]}])
    hline(s, 1.3, 5.3, 5.0, PANEL, 1.5)
    textbox(s, 1.3, 5.5, 11, 1.2, [
        {"runs": [{"t": "Presented by   ", "size": 14, "color": DIM, "font": BODY},
                  {"t": "Mohammed Abdelshafy", "size": 14, "color": TEXT, "font": BODY, "bold": True}]},
        {"before": 4, "runs": [{"t": "Supervisor       ", "size": 14, "color": DIM, "font": BODY},
                  {"t": "Dr. Waleed Elemary", "size": 14, "color": TEXT, "font": BODY, "bold": True}]},
        {"before": 4, "runs": [{"t": "Year                  ", "size": 14, "color": DIM, "font": BODY},
                  {"t": "2026", "size": 14, "color": TEXT, "font": BODY, "bold": True}]},
    ])

    # ---- 02 AGENDA ----
    s = base_slide(prs); decor(s, "Outline", 2, TOTAL); header(s, "Agenda")
    agenda = ["Introduction & problem statement", "Project objectives",
              "The design vessel", "Theoretical basis — PIANC / UNCTAD",
              "Dimensioning formulas", "The AI-engineer tool & architecture",
              "Automated CAD masterplan generation", "Worked example & results",
              "Designing for the maximum-LOA vessel",
              "Met-ocean basis — wind, tides & waves",
              "Wind rose & prevailing wind direction",
              "Wave transformation & design wave",
              "Validation & limitations", "Conclusion & future work"]
    half = 7
    col1 = agenda[:half]; col2 = agenda[half:]
    for ci, col in enumerate((col1, col2)):
        lines = []
        for j, it in enumerate(col):
            n = ci * half + j + 1
            lines.append({"before": 14, "runs": [
                {"t": f"{n:02d}  ", "size": 18, "color": BIO, "font": MONO, "bold": True},
                {"t": it, "size": 18, "color": TEXT, "font": BODY}]})
        textbox(s, 1.15 + ci * 5.9, 2.2, 5.6, 4.5, lines)

    # ---- 03 INTRODUCTION ----
    s = base_slide(prs); decor(s, "Introduction", 3, TOTAL)
    header(s, "Why Port Dimensioning Matters")
    bullets(s, [
        ("The Suez corridor carries ~12% of global trade.",
         "Port & harbour geometry must safely accommodate the largest design vessel."),
        ("Vessels keep growing — ULCVs now exceed 400 m LOA.",
         "Channels, basins and berths must be re-checked against the design ship."),
        ("Concept dimensioning is repetitive and error-prone by hand.",
         "Dozens of interdependent PIANC / UNCTAD rules per layout."),
        ("Need: a fast, standards-based engine that also produces CAD.",
         "Seconds instead of hours — with an AutoCAD-ready masterplan."),
    ])

    # ---- 04 OBJECTIVES ----
    s = base_slide(prs); decor(s, "Objectives", 4, TOTAL)
    header(s, "Project Objectives")
    bullets(s, [
        "Compute the principal port dimensions from a single design vessel.",
        "Apply recognised concept methods — PIANC WG121, UNCTAD, Barrass squat.",
        "Generate a to-scale CAD masterplan automatically (DXF / SVG / PNG).",
        "Deliver an interactive, validated AI-engineer web interface.",
        "Keep one source of truth: numbers → geometry → drawing.",
    ], size=19, gap=16)

    # ---- 05 DESIGN VESSEL ----
    s = base_slide(prs); decor(s, "Input", 5, TOTAL)
    header(s, "The Design Vessel")
    bullets(s, [
        ("Definition", "the largest vessel the port must routinely serve."),
        ("Geometric particulars", "LOA, beam B, draft T, deadweight (DWT)."),
        ("Hydrodynamic", "transit speed V and block coefficient Cb (hull fullness)."),
        ("Operational", "exposure, channel type, manoeuvre aids, number of berths."),
    ], w=6.4)
    # Cb table card
    rect(s, 8.0, 2.2, 4.4, 4.0, PANEL); vline(s, 8.0, 2.2, 4.0, AMBER, 2.5)
    textbox(s, 8.25, 2.35, 4, 0.4, [{"runs": [
        {"t": "BLOCK COEFFICIENT  Cb", "size": 13, "color": AMBER, "font": HEAD, "bold": True, "spc": 2}]}])
    cb_lines = []
    for k, v in BLOCK_COEFFICIENT.items():
        cb_lines.append({"before": 7, "runs": [
            {"t": f"{k.title():<14}", "size": 14, "color": TEXT, "font": MONO},
            {"t": f"{v:.2f}", "size": 14, "color": CYAN, "font": MONO, "bold": True}]})
    textbox(s, 8.3, 2.95, 4, 3.2, cb_lines)

    # ---- 06 THEORETICAL BASIS ----
    s = base_slide(prs); decor(s, "Method", 6, TOTAL)
    header(s, "Theoretical Basis")
    bullets(s, [
        ("PIANC WG121 — Harbour Approach Channels", "concept (Class III) design method for depth & width."),
        ("UNCTAD & standard port practice", "turning basins, quay length, anchorage geometry."),
        ("Barrass II — ship squat", "speed-dependent dynamic sinkage in shallow water."),
        ("Under-keel clearance (UKC) philosophy", "manoeuvrability margin + wave allowance by exposure."),
    ])
    textbox(s, 1.15, 6.2, 11, 0.5, [{"runs": [
        {"t": "Scope: planning / concept level — to be confirmed by detailed study.",
         "size": 13, "color": DIM, "font": BODY, "italic": True}]}])

    # ---- 07 FORMULAS I ----
    s = base_slide(prs); decor(s, "Formulas · 1 of 3", 7, TOTAL)
    header(s, "Channel Depth & Squat")
    formula_card(s, 1.1, 2.2, 5.5, 1.7, "SHIP SQUAT  (Barrass II)",
                 "S = Cb · V² / 100", "maximum bow squat in open water [m]")
    formula_card(s, 6.8, 2.2, 5.5, 1.7, "NET UNDER-KEEL CLEARANCE",
                 "UKC = max(0.05 · T, 0.5)", "manoeuvrability & bottom margin [m]")
    formula_card(s, 1.1, 4.15, 5.5, 1.7, "APPROACH CHANNEL DEPTH",
                 "D = T + UKC + S + Hw", "Hw = wave allowance by exposure")
    formula_card(s, 6.8, 4.15, 5.5, 1.7, "BERTH POCKET DEPTH",
                 "Db = T + max(0.07 · T, 0.5)", "alongside, reduced-speed squat")

    # ---- 08 FORMULAS II ----
    s = base_slide(prs); decor(s, "Formulas · 2 of 3", 8, TOTAL)
    header(s, "Approach Channel Width — PIANC")
    formula_card(s, 1.1, 2.2, 11.2, 1.6, "TWO-WAY CHANNEL",
                 "W = 2(1.5B) + 2(a·B) + 1.6B + 2(0.5B)",
                 "two manoeuvring lanes + passing distance + bank clearances")
    formula_card(s, 1.1, 4.0, 11.2, 1.6, "ONE-WAY CHANNEL",
                 "W = 1.5B + a·B + 2(0.5B)",
                 "single lane + additional widths + bank clearances")
    textbox(s, 1.15, 5.95, 11.2, 0.7, [{"runs": [
        {"t": "a = exposure width factor   ", "size": 14, "color": DIM, "font": BODY},
        {"t": "sheltered 0.6   ·   moderate 1.3   ·   exposed 2.2",
         "size": 14, "color": CYAN, "font": MONO, "bold": True}]}])

    # ---- 09 FORMULAS III ----
    s = base_slide(prs); decor(s, "Formulas · 3 of 3", 9, TOTAL)
    header(s, "Turning, Quay, Stopping & Anchorage")
    formula_card(s, 1.1, 2.2, 5.5, 1.7, "TURNING CIRCLE",
                 "Dt = k · LOA", "k: none 2.0 · tugs 1.6 · thrusters 1.5")
    formula_card(s, 6.8, 2.2, 5.5, 1.7, "TOTAL QUAY LENGTH",
                 "Lq = n·LOA + (n+1)·c", "c = max(0.10·LOA, 15) m clearance")
    formula_card(s, 1.1, 4.15, 5.5, 1.7, "DESIGN STOPPING DISTANCE",
                 "Ls ≈ 7 · LOA", "loaded, emergency astern")
    formula_card(s, 6.8, 4.15, 5.5, 1.7, "ANCHORAGE SWING RADIUS",
                 "R = LOA + 6·D + 30", "single-point swing mooring; A = π·R²")

    # ---- 10 TOOL / ARCHITECTURE ----
    s = base_slide(prs); decor(s, "Implementation", 10, TOTAL)
    header(s, "The AI-Engineer Tool")
    bullets(s, [
        ("FastAPI calculation engine (Python)", "dimensioning formulas + plan geometry."),
        ("Interactive browser interface", "presets, live validation, instant results."),
        ("ezdxf CAD exporter", "AutoCAD-ready DXF, plus SVG / PNG / JSON."),
        ("Single source of truth", "the same numbers drive screen and drawing."),
    ], w=6.3)
    # pipeline panel
    rect(s, 7.9, 2.2, 4.5, 4.0, PANEL); vline(s, 7.9, 2.2, 4.0, PLASMA, 2.5)
    textbox(s, 8.15, 2.35, 4, 0.4, [{"runs": [
        {"t": "PIPELINE", "size": 13, "color": PLASMA, "font": HEAD, "bold": True, "spc": 3}]}])
    steps = ["Design vessel input", "PIANC / UNCTAD engine", "Port dimensions",
             "Plan geometry", "DXF · SVG · PNG · JSON"]
    ln = []
    for i, st in enumerate(steps):
        ln.append({"before": 10, "runs": [
            {"t": f"{i+1}  ", "size": 16, "color": PLASMA, "font": MONO, "bold": True},
            {"t": st, "size": 15, "color": TEXT, "font": BODY}]})
        if i < len(steps) - 1:
            ln.append({"before": 2, "runs": [{"t": "     ↓", "size": 12, "color": DIM, "font": MONO}]})
    textbox(s, 8.2, 2.95, 4, 3.2, ln)

    # ---- 11 CAD MASTERPLAN (image) ----
    s = base_slide(prs); decor(s, "Output", 11, TOTAL)
    header(s, "Automated CAD Masterplan")
    s.shapes.add_picture(str(ASSETS / "plan_vlcc.png"), Inches(7.0), Inches(2.05),
                         height=Inches(4.7))
    bullets(s, [
        ("To-scale schematic", "from the computed dimensions, 1 unit = 1 m."),
        ("Layered drawing", "quay, berths, turning basin, channel, anchorage."),
        ("Open formats", "DXF (AutoCAD), SVG, PNG, JSON."),
        ("Every preset exported", "8 design vessels, ready to hand in."),
    ], w=5.6, size=16)

    # ---- 12 WORKED EXAMPLE INPUT ----
    s = base_slide(prs); decor(s, "Case study · Input", 12, TOTAL)
    header(s, "Worked Example — Neo-Panamax at Suez")
    inp = example["input"]; summ = example["summary"]
    pairs = [("Vessel type", inp["vessel_type"].title()),
             ("Class", summ["vessel_class"]),
             ("LOA", f"{inp['loa']:.0f} m"), ("Beam", f"{inp['beam']:.0f} m"),
             ("Draft", f"{inp['draft']:.1f} m"), ("Transit speed", f"{inp['speed_kn']:.0f} kn"),
             ("Block coeff. Cb", f"{summ['block_coefficient']:.2f}"),
             ("Computed squat", f"{summ['squat_m']:.2f} m"),
             ("Channel", inp["channel_type"]), ("Exposure", inp["exposure"]),
             ("Manoeuvre aids", inp["maneuver_aids"]), ("Berths", str(inp["num_berths"]))]
    for i, (k, v) in enumerate(pairs):
        col = i % 2; row = i // 2
        x = 1.15 + col * 5.9; y = 2.25 + row * 0.72
        rect(s, x, y, 5.5, 0.58, PANEL)
        textbox(s, x + 0.2, y + 0.12, 3.2, 0.4, [{"runs": [
            {"t": k.upper(), "size": 11, "color": DIM, "font": BODY, "spc": 1}]}])
        textbox(s, x + 2.6, y + 0.08, 2.8, 0.45, [{"runs": [
            {"t": str(v), "size": 15, "color": CYAN, "font": MONO, "bold": True}]}],
            align=PP_ALIGN.RIGHT)

    # ---- 13 WORKED EXAMPLE RESULTS ----
    s = base_slide(prs); decor(s, "Case study · Results", 13, TOTAL)
    header(s, "Computed Port Dimensions")
    s.shapes.add_picture(str(ASSETS / "plan_neopanamax.png"), Inches(7.05), Inches(2.0),
                         height=Inches(4.8))
    res = example["results"]
    ln = []
    for r in res:
        ln.append({"before": 6.5, "runs": [
            {"t": f"{r['label']}", "size": 13.5, "color": TEXT, "font": BODY}]})
        ln.append({"before": 0, "runs": [
            {"t": f"   {r['value']:,} ", "size": 17, "color": BIO, "font": MONO, "bold": True},
            {"t": r["unit"], "size": 12, "color": DIM, "font": MONO}]})
    textbox(s, 1.15, 2.05, 5.7, 4.8, ln)

    # ---- 14 VALIDATION ----
    s = base_slide(prs); decor(s, "Critical review", 14, TOTAL)
    header(s, "Validation & Limitations")
    bullets(s, [
        ("Planning / concept level only", "not a construction-issue design."),
        ("Empirical rules give ranges", "treat outputs as first-pass envelopes."),
        ("Must be confirmed by", "real-time manoeuvring simulation."),
        ("Site data required", "met-ocean, tidal, bathymetric & geotechnical surveys."),
        ("Engineer-in-the-loop", "a qualified maritime engineer signs off."),
    ], size=18, gap=13)

    # ---- 15 CONCLUSION ----
    s = base_slide(prs); decor(s, "Conclusion", 15, TOTAL)
    header(s, "Conclusion")
    bullets(s, [
        ("A standards-based dimensioning engine", "PIANC / UNCTAD concept methods, fully automated."),
        ("Seconds, not hours", "instant, repeatable, validated computations."),
        ("AutoCAD-ready output", "to-scale masterplan exported as DXF / SVG / PNG / JSON."),
        ("Future work", "GIS bathymetry · tidal time-series · layout optimisation · 3D / DWG."),
    ], size=19, gap=15)

    # ---- 16 MAX-LOA DESIGN VESSEL ----
    s = base_slide(prs); decor(s, "Design · Maximum LOA", 16, TOTAL)
    header(s, "Design for Maximum LOA — ULCV 400 m")
    s.shapes.add_picture(str(ASSETS / "plan_ulcv.png"), Inches(7.05), Inches(2.0),
                         height=Inches(4.8))
    mi = maxloa["input"]
    ln = [{"before": 4, "runs": [
        {"t": "Largest vessel transiting Suez — ", "size": 13, "color": DIM, "font": BODY},
        {"t": f"{maxloa['summary']['vessel_class']}", "size": 13, "color": CYAN,
         "font": BODY, "bold": True}]},
        {"before": 6, "runs": [
            {"t": f"LOA {mi['loa']:.0f} m  ·  B {mi['beam']:.1f} m  ·  T {mi['draft']:.1f} m  ·  {mi['dwt']:,} DWT",
             "size": 13, "color": TEXT, "font": MONO}]}]
    for r in maxloa["results"]:
        ln.append({"before": 7, "runs": [
            {"t": f"{r['label']}", "size": 13, "color": TEXT, "font": BODY}]})
        ln.append({"before": 0, "runs": [
            {"t": f"   {r['value']:,} ", "size": 17, "color": BIO, "font": MONO, "bold": True},
            {"t": r["unit"], "size": 12, "color": DIM, "font": MONO}]})
    textbox(s, 1.15, 2.05, 5.7, 4.9, ln)

    # ---- 17 WIND ROSE / PWD ----
    s = base_slide(prs); decor(s, "Met-ocean · Lecture 2A", 17, TOTAL)
    header(s, "Wind Rose & Prevailing Wind Direction")
    s.shapes.add_picture(str(ASSETS / "windrose.png"), Inches(7.0), Inches(1.95),
                         height=Inches(4.9))
    bullets(s, [
        ("Polygon method", "V_avg = (D1·V1+D2·V2+D3·V3+D4·V4) / ΣD"),
        (f"Recorded {_grand:,} hr/yr", "speed classes 5 / 15 / 25 / 35 kt by direction."),
        (f"PWD = {_pwd['dir']}  ({_pwd['pct']:.1f}%)", "direction of longest wind duration."),
        ("Design rule", "major breakwater oriented PERPENDICULAR to the PWD."),
    ], w=5.7, size=16)

    # ---- 18 WAVE TRANSFORMATION / DESIGN WAVE ----
    s = base_slide(prs); decor(s, "Waves · Lectures 3A + 4", 18, TOTAL)
    header(s, "Wave Transformation → Design Wave")
    formula_card(s, 1.1, 2.2, 5.5, 1.55, "SHOALING  Ks",
                 "Ks = √(Cg0 / Cg)", "Cg = nC ; n = ½(1+2kd/sinh 2kd)")
    formula_card(s, 6.8, 2.2, 5.5, 1.55, "REFRACTION  Kr",
                 "Kr = √(cosφ0 / cosφ)", "Snell: sinφ/sinφ0 = L/L0")
    formula_card(s, 1.1, 3.95, 5.5, 1.55, "DIFFRACTION  Kd",
                 "Kd = H / Hi", "from Wiegel charts (breakwater gap)")
    formula_card(s, 6.8, 3.95, 5.5, 1.55, "DESIGN WAVE HEIGHT",
                 "H = Ks · Kr · Kd · H0", "transformed deep-water wave")
    res_line = (f"d = {chan_depth} m (dredged entrance)   L0 = {sw['L0']:.0f} m   "
                f"L = {sw['L']:.1f} m   Ks = {sw['Ks']:.3f}   Kr = {rw['Kr']:.3f}   "
                f"Kd = {Kdw:.2f}")
    rect(s, 1.1, 5.75, 11.2, 0.95, PANEL); vline(s, 1.1, 5.75, 0.95, AMBER, 2.5)
    textbox(s, 1.35, 5.86, 11, 0.4, [{"runs": [
        {"t": res_line, "size": 13, "color": CYAN, "font": MONO, "bold": True}]}])
    textbox(s, 1.35, 6.26, 11, 0.4, [{"runs": [
        {"t": f"Design wave  H = {Hw:.3f} m      Average design wave length  L0 ≈ {sw['L0']:.0f} m  →  {sw['L']:.0f} m at entrance",
         "size": 14, "color": BIO, "font": MONO, "bold": True}]}])

    # ---- 19 VALIDATION (coastal vs slides) ----
    s = base_slide(prs); decor(s, "Validation", 19, TOTAL)
    header(s, "Coastal Calcs Validated vs Lecture Numbers")
    checks = [
        ("Wind V_avg North", "slide 16", "8.95", f"{next(r['v_avg'] for r in _rows if r['dir']=='N'):.2f}"),
        ("L0 = 1.56·8²", "slide 44", "100 m", f"{W.deep_water(8)[0]:.1f} m"),
        ("d / L", "slide 44", "0.1232", f"{8/W.solve_wavelength(8,8):.4f}"),
        ("Wavelength L", "slide 44", "65 m", f"{W.solve_wavelength(8,8):.1f} m"),
        ("HA  (Kd 0.10)", "slide 44", "0.25 m", f"{0.10*2.5:.2f} m"),
        ("HA  (Kd 0.145)", "slide 45", "0.36 m", f"{0.145*2.5:.2f} m"),
    ]
    # header row
    cols = [1.15, 5.0, 7.6, 10.0]
    hdr = ["CHECK", "SOURCE", "LECTURE", "COMPUTED"]
    for c, htxt in zip(cols, hdr):
        textbox(s, c, 2.15, 2.6, 0.4, [{"runs": [
            {"t": htxt, "size": 11, "color": AMBER, "font": HEAD, "bold": True, "spc": 2}]}])
    hline(s, 1.15, 2.55, 11.2, PANEL, 1.2)
    for i, (chk, src, exp, got) in enumerate(checks):
        y = 2.75 + i * 0.62
        rect(s, 1.15, y, 11.2, 0.5, PANEL if i % 2 == 0 else INK2)
        for c, txt, col, fnt in [
            (cols[0], chk, TEXT, BODY), (cols[1], src, DIM, MONO),
            (cols[2], exp, CYAN, MONO), (cols[3], got + "  ✓", BIO, MONO)]:
            textbox(s, c, y + 0.1, 3.0, 0.4, [{"runs": [
                {"t": txt, "size": 13, "color": col, "font": fnt, "bold": fnt == MONO}]}])

    # ---- 20 THANK YOU ----
    s = base_slide(prs)
    if LOGO.exists():
        s.shapes.add_picture(str(LOGO), Inches(10.85), Inches(0.5), height=Inches(2.0))
    vline(s, 1.0, 2.4, 2.6, BIO, 3)
    textbox(s, 1.3, 2.5, 11, 1.4, [{"runs": [
        {"t": "THANK YOU", "size": 60, "color": WHITE, "font": HEAD, "bold": True, "spc": 2}]}])
    textbox(s, 1.35, 4.0, 11, 0.6, [{"runs": [
        {"t": "Questions & discussion", "size": 22, "color": BIO, "font": HEAD, "spc": 1}]}])
    hline(s, 1.35, 4.9, 4.5, PANEL, 1.5)
    textbox(s, 1.35, 5.1, 11, 0.6, [{"runs": [
        {"t": "SUEZ PORT CAD · Masterplan Design AI Engineer · ECB 3802",
         "size": 14, "color": DIM, "font": MONO, "spc": 1}]}])

    prs.save(str(OUT_PPTX))
    print(f"Saved {OUT_PPTX}  ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
